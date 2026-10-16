"""Generates the fictional Northstar Plumbing & Heating mock dataset (all names and numbers are made up).
Usage: python generate_northstar.py <output_dir>   (needs openpyxl, reportlab, python-docx, pillow)"""
import csv, json, sys, os
from decimal import Decimal as D
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
from PIL import Image, ImageDraw, ImageFont

out = sys.argv[1] if len(sys.argv) > 1 else 'northstar'
os.makedirs(out, exist_ok=True)
M = ['Oct 2025','Nov 2025','Dec 2025','Jan 2026','Feb 2026','Mar 2026','Apr 2026','May 2026','Jun 2026','Jul 2026','Aug 2026','Sep 2026']
ISO = ['2025-10','2025-11','2025-12','2026-01','2026-02','2026-03','2026-04','2026-05','2026-06','2026-07','2026-08','2026-09']
q = lambda x: D(str(x)).quantize(D('0.01'))
manifest = {'business': {'name': 'Northstar Plumbing & Heating', 'industry': 'Plumbing & HVAC contractor', 'location': 'St. Cloud, MN (fictional)'}, 'planted': {}}

# ---------- 1. P&L (wide, title rows, like a QuickBooks export) ----------
service = [61200,66800,72400,74100,70300,58600,44100,58800,59300,47600,61900,62200]
install = [38400,41200,45900,43800,40100,33500,22800,27400,31600,23900,29800,34100]
maint   = [4200]*12
materials=[29800,32400,35900,35700,33100,28300,21900,24100,26400,23600,26100,28400]
subs    = [6200,6800,7400,7100,6900,5900,5200,5600,6100,5400,5800,6300]
wages   = [41800,42600,44100,44900,43800,41200,39800,40600,41900,42800,43100,43600]
fuel    = [1850,1920,2010,2090,2180,2260,2390,2510,2640,2790,2950,3150]
vehicle = [1400,1250,1600,1350,1500,1450,1550,1500,1650,1400,1600,1500]
insur   = [3100]*12
software= [640,640,640,690,690,740,790,840,890,940,990,1040]
advert  = [2400,2400,2600,2600,2500,2500,2800,2800,2700,2700,2600,2600]
rent    = [4500]*12
merchant= [2890,3110,3380,3390,3200,2690,1960,2280,2500,2080,2410,2690]
util    = [980,1150,1420,1510,1440,1210,950,870,820,840,860,910]
office  = [310,280,350,290,300,320,280,310,300,290,330,310]
inc_lines=[('Service Revenue',service),('Installation Revenue',install),('Maintenance Agreements',maint)]
cogs_lines=[('Materials & Parts',materials),('Subcontractors',subs)]
exp_lines=[('Wages & Payroll Taxes',wages),('Fuel',fuel),('Vehicle Repairs',vehicle),('Insurance',insur),('Software Subscriptions',software),('Advertising',advert),('Rent',rent),('Merchant Processing Fees',merchant),('Utilities',util),('Office Supplies',office)]
tot = lambda lines,i: sum(l[1][i] for l in lines)
ti=[tot(inc_lines,i) for i in range(12)]; tc=[tot(cogs_lines,i) for i in range(12)]; gp=[ti[i]-tc[i] for i in range(12)]
te=[tot(exp_lines,i) for i in range(12)]; ni=[gp[i]-te[i] for i in range(12)]
wb=Workbook(); ws=wb.active; ws.title='Profit and Loss'
ws.append(['Northstar Plumbing & Heating LLC']); ws.append(['Profit and Loss']); ws.append(['October 2025 - September 2026']); ws.append([])
ws.append(['Account']+M+['Total']); ws['A1'].font=Font(bold=True,size=14)
def row(name,vals): ws.append([name]+[float(v) for v in vals]+[float(sum(vals))])
ws.append(['Income']); [row(n,v) for n,v in inc_lines]; row('Total Income',ti)
ws.append(['Cost of Goods Sold']); [row(n,v) for n,v in cogs_lines]; row('Total Cost of Goods Sold',tc); row('Gross Profit',gp)
ws.append(['Expenses']); [row(n,v) for n,v in exp_lines]; row('Total Expenses',te); row('Net Income',ni)
ws.column_dimensions['A'].width=30
wb.save(f'{out}/northstar_profit_and_loss_oct2025-sep2026.xlsx')
losing=[{'month':ISO[i],'net_income':ni[i]} for i in range(12) if ni[i]<0]
manifest['planted']['losing_months']={'months':losing,'total_loss':sum(-x['net_income'] for x in losing)}
manifest['planted']['cost_trending_up']={'Fuel':{'first3':fuel[:3],'last3':fuel[-3:]},'Software Subscriptions':{'first3':software[:3],'last3':software[-3:]}}
manifest['pnl_totals']={'total_income':sum(ti),'net_income':sum(ni)}

# ---------- 2. Bank statement PDF (September 2026) ----------
bank=[('09/01','OPENING BALANCE',None),('09/02','STRIPE TRANSFER ST-8K2L',D('11842.16')),('09/02','HARTFORD INS PREMIUM',D('-3100.00')),('09/03','RENT - GRANITE PROPERTIES',D('-4500.00')),
('09/04','MOBILE DEPOSIT CHK 4471',D('6250.00')),('09/05','GUSTO PAYROLL 0905',D('-14862.40')),('09/08','FERGUSON SUPPLY #4471 ACH',D('-2418.37')),('09/09','FERGUSON SUPPLY #4471 ACH',D('-2418.37')),
('09/10','XCEL ENERGY',D('-912.44')),('09/11','STRIPE TRANSFER ST-9P4M',D('13209.55')),('09/12','JOHNSTONE SUPPLY ACH',D('-3842.10')),('09/15','MOBILE DEPOSIT CHK 4502',D('4180.00')),
('09/16','IRS USATAXPYMT',D('-5210.88')),('09/18','STRIPE TRANSFER ST-1Q7R',D('12687.30')),('09/19','GUSTO PAYROLL 0919',D('-15204.72')),('09/22','MN DEPT REVENUE',D('-1388.15')),
('09/23','WINSUPPLY ST CLOUD',D('-2975.64')),('09/25','STRIPE TRANSFER ST-6T3W',D('10944.81')),('09/26','CHASE CARD AUTOPAY',D('-8402.19')),('09/29','QUICKFLEET GPS',D('-89.00')),('09/30','SERVICE FEE',D('-25.00'))]
bal=D('48210.55'); lines=[]
for d,desc,amt in bank:
    if amt is None: lines.append((d,desc,'',f'{bal:,.2f}')); continue
    bal+=amt; lines.append((d,desc,f'{amt:,.2f}',f'{bal:,.2f}'))
c=canvas.Canvas(f'{out}/northstar_bank_statement_sep2026.pdf',pagesize=letter)
c.setFont('Helvetica-Bold',13); c.drawString(50,750,'North Country Community Bank'); c.setFont('Helvetica',10)
c.drawString(50,734,'Business Checking Statement   Account ending 4471'); c.drawString(50,720,'Northstar Plumbing & Heating LLC   Statement period 09/01/2026 - 09/30/2026')
c.setFont('Helvetica-Bold',10); y=694
for x,h in [(50,'Date'),(110,'Description'),(390,'Amount'),(480,'Balance')]: c.drawString(x,y,h)
c.setFont('Helvetica',10); y-=18
for d,desc,amt,b in lines:
    c.drawString(50,y,d); c.drawString(110,y,desc); c.drawRightString(450,y,amt); c.drawRightString(550,y,b); y-=16
c.drawString(50,y-14,f'Closing balance {bal:,.2f}'); c.save()
manifest['planted']['duplicate_bank_charge']={'vendor':'FERGUSON SUPPLY #4471 ACH','amount':2418.37,'dates':['2026-09-08','2026-09-09']}

# ---------- 3. Credit card CSV (Apr-Sep 2026; negative = charge) ----------
card=[]
fieldpro=[149,149,199,199,249,249]
for k,mon in enumerate(['04','05','06','07','08','09']):
    card+= [(f'2026-{mon}-03','FIELDPRO SOFTWARE SUBSCRIPTION','Software',-fieldpro[k]),(f'2026-{mon}-05','QUICKFLEET GPS MONTHLY','Software',-89.00),(f'2026-{mon}-05','FLEETTRACK PRO GPS','Software',-94.00),
            (f'2026-{mon}-07','KWIK TRIP #612 ST CLOUD','Gas',-(212.40+k*18.5)),(f'2026-{mon}-14','HOLIDAY STATIONSTORE 302','Gas',-(188.75+k*21.25)),(f'2026-{mon}-17','MENARDS ST CLOUD','Home Improvement',-(533.10+k*11)),
            (f'2026-{mon}-21','VERIZON WIRELESS','Bills & Utilities',-412.86),(f'2026-{mon}-24','GOOGLE ADS','Advertising',-(300.00 if k%2 else 350.00))]
card.append(('2026-07-14','HOME DEPOT #2231','Home Improvement',-612.48)); card.append(('2026-07-14','HOME DEPOT #2231','Home Improvement',-612.48))
card.append(('2026-08-11','CULVERS ST CLOUD','Food & Drink',-46.22))
card.sort(key=lambda r:r[0])
with open(f'{out}/northstar_credit_card_apr-sep2026.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['Transaction Date','Post Date','Description','Category','Type','Amount'])
    for d,desc,cat,a in card: w.writerow([d,d,desc,cat,'Sale',f'{q(a)}'])
manifest['planted']['subscription_creep']={'vendor':'FIELDPRO SOFTWARE SUBSCRIPTION','monthly':fieldpro,'excess_over_first':sum(x-fieldpro[0] for x in fieldpro)}
manifest['planted']['duplicate_card_charge']={'vendor':'HOME DEPOT #2231','amount':612.48,'date':'2026-07-14'}
manifest['planted']['overlapping_subscriptions_grok_only']=['QUICKFLEET GPS MONTHLY 89.00','FLEETTRACK PRO GPS 94.00']

# ---------- 4. Invoices / AR (CSV) ----------
inv=[('INV-2036','Riverside Clinic','2026-07-06','2026-08-05',3420,3420,'Paid'),('INV-2037','M. Schultz','2026-07-09','2026-08-08',860,860,'Paid'),('INV-2038','Sauk Rapids Schools','2026-07-12','2026-08-11',9150,9150,'Paid'),
('INV-2039','K. Nguyen','2026-07-15','2026-08-14',1240,1240,'Paid'),('INV-2040','Benton Ag Supply','2026-07-18','2026-08-17',2780,2780,'Paid'),('INV-2041','Lakeside Dental','2026-07-16','2026-08-15',4850,0,'Open'),
('INV-2042','T. Becker','2026-07-22','2026-08-21',695,695,'Paid'),('INV-2043','Waite Park Hotel','2026-07-28','2026-08-27',5610,5610,'Paid'),('INV-2044','J. Larson','2026-08-02','2026-09-01',1385,1385,'Paid'),
('INV-2045','Northway Church','2026-08-05','2026-09-04',2240,2240,'Paid'),('INV-2046','D. Patel','2026-08-08','2026-09-07',930,930,'Paid'),('INV-2047','Granite City Apartments','2026-07-31','2026-08-30',7325,2000,'Partial'),
('INV-2048','Sartell Foods','2026-08-14','2026-09-13',3175,3175,'Paid'),('INV-2049','A. Johnson','2026-08-19','2026-09-18',540,540,'Paid'),('INV-2050','Crossroads Mall','2026-08-22','2026-09-21',6480,6480,'Paid'),
('INV-2051','L. Meyer','2026-08-25','2026-09-24',1120,1120,'Paid'),('INV-2052','R. Hoffman','2026-08-11','2026-09-10',1280,0,'Open'),('INV-2053','St. Joseph Bakery','2026-09-01','2026-10-01',2365,2365,'Paid'),
('INV-2054','E. Kowalski','2026-09-04','2026-10-04',780,780,'Paid'),('INV-2055','Rice Veterinary','2026-09-09','2026-10-09',3940,3940,'Paid'),('INV-2056','P. Anderson','2026-09-15','2026-10-30',1460,0,'Open'),
('INV-2057','Clearwater Motel','2026-09-22','2026-10-22',4210,0,'Open'),('INV-2058','St. Cloud Brewing Co','2026-08-26','2026-09-25',3960,0,'Open'),('INV-2059','B. Thompson','2026-09-28','2026-10-28',990,0,'Open')]
with open(f'{out}/northstar_invoices_ar_jul-sep2026.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['Invoice #','Customer','Invoice Date','Due Date','Amount','Amount Paid','Status'])
    for r in inv: w.writerow([r[0],r[1],r[2],r[3],f'{r[4]:.2f}',f'{r[5]:.2f}',r[6]])
asof='2026-10-16'; od=[{'invoice':r[0],'outstanding':r[4]-r[5]} for r in inv if r[3]<asof and r[5]<r[4]]
manifest['planted']['overdue_invoices']={'as_of':asof,'invoices':od,'total':sum(x['outstanding'] for x in od)}

# ---------- 5. Job costing (CSV) ----------
# Field techs (Lisa Hagen is the office manager). Sam Voigt is assigned the thin-margin water heater installs.
FIELD=['Jake Lindgren','Marcus Olson','Ana Reyes','Tom Brekke','Sam Voigt']
SAM_JOBS={'J-3110','J-3115','J-3121','J-3124'}
def tech_of(jid):
    if jid in SAM_JOBS: return 'Sam Voigt'
    return FIELD[int(jid[2:])%4]
SOURCES=['Repeat customer','Google Ads','Referral','Google Ads','Repeat customer','Facebook Ads']
ANGI={'J-3103','J-3117'}
def src_of(jid):
    return 'Angi Leads' if jid in ANGI else SOURCES[int(jid[2:])%6]
CALLBACKS={'J-3121','J-3126'}

jobs=[]; n=3101
base=[('Drain cleaning',480,150,40),('Furnace repair',920,310,220),('AC tune-up',260,90,25),('Boiler service',740,260,120),('Water heater install',1850,610,780),('Leak repair',540,180,70)]
plan={3107:('Boiler replacement',6200,3950,3180),3110:('Water heater install',1720,690,890),3115:('Water heater install',1790,720,910),3121:('Water heater install',1650,980,1140),3124:('Water heater install',1760,700,905),3128:('Repipe',9800,6450,4120)}
days=['07-02','07-06','07-09','07-13','07-16','07-20','07-23','07-27','07-30','08-03','08-06','08-10','08-13','08-17','08-20','08-24','08-27','08-31','09-03','09-07','09-10','09-14','09-17','09-21','09-24','09-28','09-29','09-30','09-30','09-30']
cust=['Riverside Clinic','M. Schultz','K. Nguyen','Benton Ag Supply','T. Becker','Waite Park Hotel','J. Larson','Northway Church','D. Patel','Granite City Apts','Sartell Foods','A. Johnson','Crossroads Mall','L. Meyer','R. Hoffman','St. Joseph Bakery','E. Kowalski','Rice Veterinary','P. Anderson','Clearwater Motel','St. Cloud Brewing','B. Thompson','Lakeside Dental','Sauk Rapids Schools','C. Olson','G. Weber','H. Klein','I. Ruiz','N. Peters','O. Hanson']
for i in range(30):
    jid=n+i
    if jid in plan: t,r,l,m=plan[jid]
    else: t,r,l,m=base[i%6]; r=r+(i*15); l=l+(i*4); m=m+(i*3)
    jobs.append((f'J-{jid}',f'2026-{days[i]}',cust[i],t,r,l,m))
with open(f'{out}/northstar_job_costing_jul-sep2026.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['Job #','Completed Date','Customer','Job Type','Technician','Lead Source','Callback','Labor Hours','Revenue','Labor Cost','Material Cost'])
    for j in jobs: w.writerow([j[0],j[1],j[2],j[3],tech_of(j[0]),src_of(j[0]),'Yes' if j[0] in CALLBACKS else 'No',f'{j[5]/45:.1f}',f'{j[4]:.2f}',f'{j[5]:.2f}',f'{j[6]:.2f}'])
lj=[{'job':j[0],'loss':j[5]+j[6]-j[4]} for j in jobs if j[4]-j[5]-j[6]<0]
manifest['planted']['losing_jobs']={'jobs':lj,'total':sum(x['loss'] for x in lj)}
manifest['planted']['underpriced_jobs']={'jobs':['J-3110','J-3115','J-3124'],'note':'Water heater installs with positive but thin margins'}


# ---------- 5b. Planted labor/marketing/callback expectations (same formulas as lib/leaks/engine.ts) ----------
def margin(j): return j[4]-j[5]-j[6]
techs={}
for j in jobs:
    t=techs.setdefault(tech_of(j[0]),{'rev':0,'m':0}); t['rev']+=j[4]; t['m']+=margin(j)
rates={k:v['m']/v['rev'] for k,v in techs.items()}
srt=sorted(rates.values()); med=(srt[len(srt)//2] if len(srt)%2 else (srt[len(srt)//2-1]+srt[len(srt)//2])/2)
low={k:round(techs[k]['rev']*(med-r)*100) for k,r in rates.items() if r < med-0.10}
manifest['planted']['technician_contribution']={'median_margin':round(med,4),'below':{k:v/100 for k,v in low.items()},'total':sum(low.values())/100}
cb=[{'job':j[0],'cost':j[5]+j[6]} for j in jobs if j[0] in CALLBACKS]
manifest['planted']['callbacks']={'jobs':cb,'total':sum(x['cost'] for x in cb)}
MKT={'Google Ads':(650,14),'Facebook Ads':(420,6),'Angi Leads':(780,9)}
mrows=[]
for mo,end in [('07','31'),('08','31'),('09','30')]:
    for src,(sp,ld) in MKT.items(): mrows.append((f'2026-{mo}-{end}',src,sp,ld))
with open(f'{out}/northstar_marketing_spend_jul-sep2026.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['Month End','Lead Source','Ad Spend','Leads'])
    for r in mrows: w.writerow([r[0],r[1],f'{r[2]:.2f}',r[3]])
mk={}
for src,(sp,_) in MKT.items():
    contrib=sum(margin(j) for j in jobs if src_of(j[0])==src); spend=sp*3
    if contrib<spend: mk[src]={'spend':spend,'job_contribution':contrib,'shortfall':spend-max(0,contrib)}
manifest['planted']['marketing_return']={'sources':mk,'total':sum(v['shortfall'] for v in mk.values())}

# ---------- 6. Payroll (XLSX) ----------
emps=[('Jake Lindgren',38.0),('Marcus Olson',36.5),('Ana Reyes',32.0),('Tom Brekke',30.0),('Lisa Hagen',24.0),('Sam Voigt',22.5)]
otplan={'Jake Lindgren':[18,22,20,24,19,21,23],'Marcus Olson':[16,20,22,18,21,24,20],'Ana Reyes':[2,0,3,0,2,1,0],'Tom Brekke':[0,2,0,1,0,0,2],'Lisa Hagen':[0]*7,'Sam Voigt':[1,0,0,2,0,1,0]}
paydates=['2026-07-10','2026-07-24','2026-08-07','2026-08-21','2026-09-04','2026-09-18','2026-10-02']
wb=Workbook(); ws=wb.active; ws.title='Payroll Register'
ws.append(['Northstar Plumbing & Heating LLC - Payroll Register (Jul-Sep 2026)']); ws.append([])
ws.append(['Pay Date','Employee','Regular Hours','Overtime Hours','Hourly Rate','Overtime Pay','Gross Wages'])
tot_g=D(0); tot_ot=D(0)
for k,pd in enumerate(paydates):
    for e,rate in emps:
        oh=otplan[e][k]; reg=80; otp=q(D(str(rate))*D('1.5')*oh); g=q(D(str(rate))*reg)+otp; tot_g+=g; tot_ot+=otp
        ws.append([pd,e,reg,oh,rate,float(otp),float(g)])
wb.save(f'{out}/northstar_payroll_register_jul-sep2026.xlsx')
manifest['planted']['overtime']={'total_overtime_pay':float(tot_ot),'total_gross':float(tot_g),'overtime_premium_estimate':float(q(tot_ot/3)),'heavy_ot':['Jake Lindgren','Marcus Olson']}


# ---------- 6b. Timesheets (CSV): billable vs non-billable (drive/shop) hours ----------
BILL={'Jake Lindgren':0.58,'Marcus Olson':0.61,'Ana Reyes':0.74,'Tom Brekke':0.76,'Sam Voigt':0.72}
rate_of=dict(emps); util_loss=D(0); ut={}
with open(f'{out}/northstar_timesheets_jul-sep2026.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['Period End','Employee','Billable Hours','Non-Billable Hours'])
    for k,pd in enumerate(paydates):
        for e in FIELD:
            tot=80+otplan[e][k]; b=round(tot*BILL[e],1); nb=round(tot-b,1); w.writerow([pd,e,f'{b:.1f}',f'{nb:.1f}'])
            u=ut.setdefault(e,[0.0,0.0]); u[0]+=b; u[1]+=b+nb
for e,(b,t) in ut.items():
    gap=0.70*t-b
    if gap>0: util_loss+=q(D(str(round(gap,4)))*D(str(rate_of[e])))
manifest['planted']['labor_utilization']={'target':0.70,'by_employee':{e:round(b/t,4) for e,(b,t) in ut.items()},'below_target':[e for e,(b,t) in ut.items() if b/t<0.70],'total':float(util_loss)}

# ---------- 7. Processor fees (CSV, odd export name) ----------
gross=[q(x*0.48) for x in ti]
with open(f'{out}/stripe_payouts_export_2026.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['Payout Date','Gross Volume','Fees','Net Payout','Currency'])
    tf=D(0); tg=D(0)
    for i,iso in enumerate(ISO):
        fee=q(gross[i]*D('0.0345')+D('48.00')); tf+=fee; tg+=gross[i]
        w.writerow([f'{iso}-28',f'{gross[i]}',f'{fee}',f'{gross[i]-fee}','usd'])
manifest['planted']['processor_fees']={'total_fees':float(tf),'total_gross':float(tg),'effective_rate':float(q(tf/tg*100)),'excess_over_2_9pct':float(q(tf-tg*D('0.029')))}

manifest['planted']['revenue_trend']={'first3':float(sum(ti[:3])),'last3':float(sum(ti[-3:])),'decline':float(sum(ti[:3])-sum(ti[-3:]))}

# ---------- 8. Business notes (DOCX) ----------
doc=Document(); doc.add_heading('Northstar Plumbing & Heating - Owner Notes',1)
for p in ['Owner: Erik Halvorsen (fictional). Based in St. Cloud, MN. Residential and light commercial plumbing, heating and cooling.',
 'Team: 5 field techs and 1 office manager (Lisa). 3 service vans. Busy season is heating (Oct-Feb).',
 'Pricing for water heater installs has not changed since 2023 even though tank prices went up.',
 'We switched GPS providers to FleetTrack Pro in the spring but I am not sure QuickFleet was ever cancelled.',
 'Jake and Marcus are working a lot of overtime because we have not hired a third senior tech yet.',
 'We pay for Google Ads, Facebook Ads and Angi leads; I have never checked which ones actually bring in profitable jobs.',
 'We had a couple of callbacks this summer (J-3121, J-3126) that we went back to for free.',
 'Some commercial customers (Lakeside Dental, Granite City Apartments) are slow to pay; nobody is assigned to follow up.',
 'Fuel costs feel higher every month; routes are not planned, techs drive back to the shop between calls.']: doc.add_paragraph(p)
doc.save(f'{out}/northstar_business_notes.docx')

# ---------- 9. Receipt image (for vision) ----------
img=Image.new('RGB',(520,700),'white'); d=ImageDraw.Draw(img)
try: font=ImageFont.truetype('DejaVuSansMono.ttf',22); big=ImageFont.truetype('DejaVuSansMono-Bold.ttf',28)
except Exception: font=big=ImageFont.load_default()
y=30
for t,f_ in [('MENARDS',big),('St. Cloud, MN',font),('09/15/2026  14:32',font),('',font),('PEX PIPE 1/2 X100   89.97',font),('SHARKBITE COUPLR x6  71.94',font),('WATER HEATER EXP TK 64.99',font),('COPPER FITTINGS     112.48',font),('',font),('SUBTOTAL           339.38',font),('TAX                 22.92',font),('TOTAL              362.30',big),('',font),('VISA ****4471',font)]:
    d.text((30,y),t,fill='black',font=f_); y+=40
img.save(f'{out}/receipt_menards_2026-09-15.png')
manifest['planted']['receipt_image']={'total':362.30,'note':'Image: AI-read via Grok vision, requires owner confirmation'}
json.dump(manifest,open(f'{out}/../northstar_expected_leaks.json','w'),indent=2)
print(json.dumps(manifest['planted'],indent=1))
